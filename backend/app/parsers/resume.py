import re,fitz
from docx import Document
HEADINGS={'contact':'contact','summary':'summary','objective':'objective','experience':'experience','work experience':'experience','professional experience':'experience','employment history':'experience','education':'education','skills':'skills','technical skills':'skills','core skills':'skills','projects':'projects','certifications':'certifications','achievements':'achievements','awards':'awards','publications':'publications','languages':'languages','volunteering':'volunteering'}
def sections(text):
 out={};current='other';buf=[]
 for line in text.splitlines():
  key=re.sub(r'[^a-z ]','',line.lower()).strip()
  if key in HEADINGS:
   if buf:out[current]='\n'.join(buf).strip()
   current=HEADINGS[key];buf=[]
  else:buf.append(line)
 if buf:out[current]='\n'.join(buf).strip()
 return {k:v for k,v in out.items() if v}
def parse_pdf(path):return '\n'.join(page.get_text() for page in fitz.open(path))
def parse_docx(path):
 d=Document(path);lines=[p.text for p in d.paragraphs];lines += [' | '.join(c.text for c in row.cells) for t in d.tables for row in t.rows];return '\n'.join(lines)
